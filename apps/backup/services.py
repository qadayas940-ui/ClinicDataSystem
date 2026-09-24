import csv
import hashlib
import json
import shutil
import sqlite3
import subprocess
import tempfile
import zipfile
from pathlib import Path

from django.conf import settings
from django.core.management import call_command
from django.db import connection
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from apps.core.models import BackupHistory


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _stamp():
    return timezone.localtime().strftime("%Y%m%d-%H%M%S")


def _export_dir(kind):
    path = Path(settings.DATA_PATH) / "Files" / kind
    path.mkdir(parents=True, exist_ok=True)
    return path


def _database_snapshot(destination):
    """Create a consistent snapshot for SQLite or PostgreSQL without fixed paths."""
    database = settings.DATABASES["default"]
    vendor = connection.vendor
    connection.close()
    if vendor == "sqlite":
        output = destination / "clinic.db"
        source = sqlite3.connect(database["NAME"])
        target = sqlite3.connect(output)
        try:
            source.backup(target)
        finally:
            target.close()
            source.close()
        return output, "sqlite"

    if vendor == "postgresql" and shutil.which("pg_dump"):
        output = destination / "clinic-postgresql.dump"
        command = ["pg_dump", "--format=custom", "--file", str(output)]
        for option, flag in (("HOST", "--host"), ("PORT", "--port"), ("USER", "--username")):
            if database.get(option):
                command.extend([flag, str(database[option])])
        command.append(str(database["NAME"]))
        environment = None
        if database.get("PASSWORD"):
            import os
            environment = os.environ.copy()
            environment["PGPASSWORD"] = str(database["PASSWORD"])
        subprocess.run(command, check=True, timeout=1800, env=environment, capture_output=True)
        return output, "postgresql-custom"

    output = destination / "clinic-data.json"
    call_command(
        "dumpdata", "--all", "--natural-foreign", "--natural-primary",
        "--exclude", "contenttypes", "--exclude", "auth.permission",
        output=str(output), verbosity=0,
    )
    return output, f"{vendor}-django-json"


def create_database_backup(user=None, backup_type="manual"):
    backup_dir = Path(settings.DATA_PATH) / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    final_path = backup_dir / f"ClinicData-{_stamp()}.clinicbackup"
    history = BackupHistory.objects.create(
        backup_type=backup_type, status="in_progress", created_by=user,
        app_version=settings.APP_VERSION, db_version=settings.DB_SCHEMA_VERSION,
    )
    try:
        with tempfile.TemporaryDirectory() as temporary:
            temporary_path = Path(temporary)
            db_copy, database_format = _database_snapshot(temporary_path)
            manifest = {
                "format": "ClinicDataBackup/2", "created_at": timezone.now().isoformat(),
                "app_version": settings.APP_VERSION, "db_schema_version": settings.DB_SCHEMA_VERSION,
                "database_engine": connection.vendor, "database_format": database_format,
                "database_file": db_copy.name, "database_sha256": sha256_file(db_copy),
            }
            manifest_path = temporary_path / "manifest.json"
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            with zipfile.ZipFile(final_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
                archive.write(db_copy, db_copy.name)
                archive.write(manifest_path, "manifest.json")
        history.file_path = str(final_path)
        history.file_size = final_path.stat().st_size
        history.checksum = sha256_file(final_path)
        history.status = "success"
        history.save()
        return history
    except Exception as exc:
        history.status = "failed"
        history.notes = str(exc)[:1000]
        history.save(update_fields=["status", "notes"])
        raise


def _export_tables():
    from apps.laboratory.models import LabOrder
    from apps.ophthalmology.models import EyeClinicVisit
    from apps.patients.models import Patient
    from apps.referrals.models import Referral

    patients = Patient.objects.select_related("primary_name").prefetch_related(
        "contacts", "addresses", "visits__department", "visits__doctor_reference", "visits__organizer_reference"
    ).order_by("created_at", "pk")
    patient_rows = []
    sequence = 0
    for patient in patients:
        visits = list(patient.visits.order_by("visit_date", "pk"))
        if not visits:
            visits = [None]
        for visit in visits:
            sequence += 1
            patient_rows.append([
                sequence, patient.internal_code, patient.display_name, patient.get_gender_display(),
                patient.calculated_age, patient.addresses.values_list("text", flat=True).first() or "",
                patient.contacts.values_list("value", flat=True).first() or "",
                visit.department_label if visit else "", visit.diagnosis if visit else "",
                visit.doctor_label if visit else "",
                visit.organizer_label if visit else "",
                timezone.localtime(visit.visit_date).strftime("%Y/%m/%d — %I:%M %p").lstrip("0") if visit else "", visit.notes if visit else "",
            ])

    lab_rows = []
    for order in LabOrder.objects.select_related("patient", "patient__primary_name").prefetch_related("tests"):
        tests = list(order.tests.all()) or [None]
        for test in tests:
            lab_rows.append([
                len(lab_rows) + 1, order.patient.internal_code, order.patient.display_name,
                order.patient.calculated_age, test.test_name if test else "",
            ])

    referral_rows = [
        [index, item.patient.internal_code, item.patient.display_name, item.destination_name]
        for index, item in enumerate(
            Referral.objects.select_related("patient", "patient__primary_name"), start=1
        )
    ]
    eye_rows = [
        [index, item.patient.internal_code, item.patient.display_name, item.patient.calculated_age,
         item.diagnosis, item.notes]
        for index, item in enumerate(
            EyeClinicVisit.objects.select_related("patient", "patient__primary_name"), start=1
        )
    ]
    return [
        ("المرضى", ["ت", "الرقم التعريفي الخاص بالمريض", "الاسم", "الجنس", "العمر", "العنوان",
                     "رقم الهاتف", "القسم", "الحالة", "اسم الطبيب", "اسم المنظم", "التاريخ", "الملاحظات"], patient_rows),
        ("المختبر", ["ت", "الرقم التعريفي", "الاسم", "العمر", "نوع الفحص"], lab_rows),
        ("الإحالات", ["ت", "الرقم التعريفي", "الاسم", "الطبيب المحال إليه"], referral_rows),
        ("عيادة العيون", ["ت", "الرقم التعريفي", "الاسم", "العمر", "التشخيص", "الملاحظات"], eye_rows),
    ]


def _plain(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _style_excel_sheet(sheet, table_name=None):
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F73")
    if table_name and sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def _safe_sheet_title(value, existing):
    cleaned = "".join("_" if char in '[]:*?/\\\\' else char for char in str(value or "غير محدد")).strip()[:25]
    base = f"قسم - {cleaned or 'غير محدد'}"
    title, counter = base[:31], 2
    while title in existing:
        suffix = f"-{counter}"
        title = base[:31-len(suffix)] + suffix
        counter += 1
    return title


def create_excel_export():
    path = _export_dir("Excel") / f"ClinicData-export-{_stamp()}.xlsx"
    book = Workbook()
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    book.remove(book.active)
    exported = _export_tables()
    for index, (title, headers, rows) in enumerate(exported, start=1):
        sheet = book.create_sheet(title=title)
        sheet.append(headers)
        for row in rows:
            sheet.append([_plain(value) for value in row])
        _style_excel_sheet(sheet, f"ClinicTable{index}")
    patient_headers, patient_rows = exported[0][1], exported[0][2]
    index_sheet = book.create_sheet("اختصاصات")
    index_sheet.append(["الاختصاص", "عدد السجلات", "اسم الورقة"])
    groups = {}
    for row in patient_rows:
        groups.setdefault(str(row[7] or "غير محدد"), []).append(row)
    existing = set(book.sheetnames)
    for number, (department, rows) in enumerate(sorted(groups.items()), start=1):
        title = _safe_sheet_title(department, existing)
        existing.add(title)
        specialty = book.create_sheet(title)
        specialty.append(patient_headers)
        for row in rows:
            specialty.append([_plain(value) for value in row])
        _style_excel_sheet(specialty, f"SpecialtyTable{number}")
        index_sheet.append([department, len(rows), title])
        index_sheet.cell(index_sheet.max_row, 3).hyperlink = f"#'{title}'!A1"
    _style_excel_sheet(index_sheet, "SpecialtiesIndex")
    search = book.create_sheet("البحث")
    search.sheet_view.rightToLeft = True
    search.merge_cells("A1:F1")
    search["A1"] = "البحث عن مريض"
    search["A1"].font = Font(bold=True, color="FFFFFF", size=15)
    search["A1"].fill = PatternFill("solid", fgColor="062F50")
    search["A1"].alignment = Alignment(horizontal="center", vertical="center")
    search.row_dimensions[1].height = 32

    search.merge_cells("A2:D2")
    search["A2"] = "ابحث هنا ← اكتب جزءًا من اسم المريض"
    search["A2"].font = Font(bold=True, color="062F50", size=12)
    search["A2"].alignment = Alignment(horizontal="right", vertical="center")
    search.row_dimensions[2].height = 29
    search.merge_cells("B3:D3")
    search["B3"] = ""
    search["B3"].font = Font(size=14, color="062F50")
    search["B3"].fill = PatternFill("solid", fgColor="E8F6F2")
    search["B3"].border = Border(bottom=Side(style="medium", color="087F73"))
    search["B3"].alignment = Alignment(horizontal="right", vertical="center")
    search.row_dimensions[3].height = 34
    search["A4"] = "اختر الاسم"
    search["A4"].font = Font(bold=True, color="062F50")
    search.merge_cells("B4:D4")
    search["B4"].fill = PatternFill("solid", fgColor="F1F7FA")
    search["B4"].alignment = Alignment(horizontal="right", vertical="center")
    search.row_dimensions[4].height = 30
    search.merge_cells("B6:D6")
    search["B6"] = "الأسماء المطابقة — الاسم | الرقم التعريفي"
    search["B6"].font = Font(bold=True, color="FFFFFF")
    search["B6"].fill = PatternFill("solid", fgColor="087F73")
    search["B6"].alignment = Alignment(horizontal="right")
    last_patient_row = len(patient_rows) + 1
    if patient_rows:
        names = f"المرضى!$C$2:$C${last_patient_row}"
        codes = f"المرضى!$B$2:$B${last_patient_row}"
        for row_number in range(7, 22):
            search.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=4)
            cell = search[f"B{row_number}"]
            cell.value = (
                f'=IF($B$3="","",IFERROR(INDEX(_xlfn.UNIQUE(_xlfn._xlws.FILTER('
                f'{names}&" | "&{codes},ISNUMBER(SEARCH($B$3,{names})),"")),'
                f'ROW()-6),""))'
            )
            cell.alignment = Alignment(horizontal="right")
            if row_number % 2:
                cell.fill = PatternFill("solid", fgColor="F1F7FA")
        validation = DataValidation(type="list", formula1="$B$7:$B$21", allow_blank=True)
        validation.showDropDown = False
        validation.error = "اختر اسمًا من النتائج المطابقة."
        validation.showErrorMessage = True
        search.add_data_validation(validation)
        validation.add(search["B4"])

    search.merge_cells("G2:I2")
    search["G2"] = "بطاقة المريض"
    search["G2"].font = Font(bold=True, color="FFFFFF", size=14)
    search["G2"].fill = PatternFill("solid", fgColor="062F50")
    search["G2"].alignment = Alignment(horizontal="center")
    search["G4"] = "الرقم التعريفي"
    search.merge_cells("H4:I4")
    search["H4"] = '=IFERROR(TRIM(MID($B$4,FIND(" | ",$B$4)+3,99)),"")'
    card_rows = [
        (3, "الاسم", "C", False), (5, "الجنس", "D", False),
        (6, "العمر", "E", False), (7, "العنوان", "F", False),
        (8, "الهاتف", "G", False), (9, "القسم", "H", True),
        (10, "الحالة / التشخيص", "I", True), (11, "الطبيب", "J", True),
        (12, "المنظّم", "K", True), (13, "التاريخ", "L", True),
        (14, "الملاحظات", "M", True),
    ]
    for row_number, label, source_col, latest in card_rows:
        search[f"G{row_number}"] = label
        search.merge_cells(start_row=row_number, start_column=8, end_row=row_number, end_column=9)
        if patient_rows:
            if latest:
                result = f'LOOKUP(2,1/(المرضى!$B$2:$B${last_patient_row}=$H$4),المرضى!${source_col}$2:${source_col}${last_patient_row})'
            else:
                result = f'INDEX(المرضى!${source_col}:${source_col},MATCH($H$4,المرضى!$B:$B,0))'
            search[f"H{row_number}"] = f'=IF($H$4="","",IFERROR({result},""))'
    for row_number in range(3, 15):
        search[f"G{row_number}"].font = Font(bold=True, color="062F50")
        search[f"G{row_number}"].fill = PatternFill("solid", fgColor="E8F6F2")
        search[f"H{row_number}"].alignment = Alignment(horizontal="right", vertical="center")
        search.row_dimensions[row_number].height = max(search.row_dimensions[row_number].height or 0, 27)

    search["A24"] = "جميع زيارات المريض المختار (Excel 365)"
    search["A24"].font = Font(bold=True, color="062F50")
    if patient_rows:
        search["A25"] = f'=IF($H$4="","",_xlfn._xlws.FILTER(المرضى!A2:M${last_patient_row},المرضى!B2:B${last_patient_row}=$H$4,"لا توجد زيارات"))'
    for column, width in {"A":25,"B":31,"C":20,"D":25,"E":15,"F":20,"G":22,"H":28,"I":24}.items():
        search.column_dimensions[column].width = width
    changes = book.create_sheet("التغييرات")
    changes.append(["العملية","الرقم التعريفي","الاسم","القسم","التاريخ","الملاحظات"])
    changes.append(["","","","","","تُراجع عبر شاشة الاستيراد قبل الدمج؛ لا تُستبدل قاعدة البيانات مباشرة."])
    operation_validation = DataValidation(type="list", formula1='"إضافة,تعديل,أرشفة"', allow_blank=True)
    changes.add_data_validation(operation_validation)
    operation_validation.add("A2:A5000")
    _style_excel_sheet(changes, "ChangesTable")
    book.save(path)
    return path

def create_csv_export():
    path = _export_dir("ZIP") / f"ClinicData-export-{_stamp()}.zip"
    with tempfile.TemporaryDirectory() as temporary:
        temporary_path = Path(temporary)
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for title, headers, rows in _export_tables():
                csv_path = temporary_path / f"{title}.csv"
                with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
                    writer = csv.writer(stream)
                    writer.writerow(headers)
                    writer.writerows([_plain(value) for value in row] for row in rows)
                archive.write(csv_path, csv_path.name)
    return path


def create_json_export():
    path = _export_dir("JSON") / f"ClinicData-export-{_stamp()}.json"
    payload = {}
    for title, headers, rows in _export_tables():
        payload[title] = [dict(zip(headers, [_plain(value) for value in row])) for row in rows]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
