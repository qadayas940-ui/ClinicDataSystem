import hashlib
import io
import json
import tempfile
import zipfile
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from django.test import override_settings
from openpyxl import Workbook, load_workbook

from apps.accounts.models import Role
from apps.core.models import Department, Notification, ReferenceValue
from apps.importer.models import ImportBatch, ImportSheet, SourceRow
from apps.importer.services import analyze_workbook, import_batch_records, source_repeat_count
from apps.laboratory.models import LabOrder, LabOrderTest
from apps.patients.forms import PatientForm
from apps.patients.models import Patient
from apps.patients.services import create_patient, find_patient_candidates
from apps.referrals.models import Referral
from apps.visits.models import Visit

User = get_user_model()


class PatientWorkflowTests(TestCase):
    def setUp(self):
        self.role = Role.objects.create(name="المنظم", code=Role.CODE_ORGANIZER)
        self.user = User.objects.create_user(username="org", password="StrongPass123", role=self.role)

    def test_patient_gets_permanent_code_and_derived_age(self):
        patient = create_patient({"full_name": "اختبار مريض كامل علي", "gender": "male", "date_of_birth": timezone.localdate().replace(year=timezone.localdate().year - 20), "approx_age_value": None, "approx_age_unit": "", "phone": "0770 123 4567", "address": "الموصل"}, self.user)
        self.assertRegex(patient.internal_code, rf"^CLN{str(timezone.localdate().year)[-2:]}-\d{{10}}$")
        self.assertEqual(patient.display_name, "اختبار مريض كامل علي")
        self.assertIn("20", patient.calculated_age)
        self.assertEqual(patient.contacts.first().value, "+9647701234567")

    def test_search_by_name(self):
        patient = create_patient({"full_name": "سارة أحمد محمود علي", "gender": "female", "date_of_birth": None, "approx_age_value": 8, "approx_age_unit": "year", "phone": "07899189225", "address": "الزهور"}, self.user)
        self.client.force_login(self.user)
        response = self.client.get(reverse("patients:list"), {"q": "سارة"})
        self.assertContains(response, patient.internal_code)

    def test_unified_search_and_match_find_manual_or_imported_patient(self):
        manual = create_patient({"full_name": "سارة أحمد محمود علي", "gender": "female", "date_of_birth": None, "approx_age_value": 8, "approx_age_unit": "year", "phone": "07899189225", "address": "الزهور"}, self.user)
        imported = create_patient({"full_name": "حيدر سالم كاظم حسن", "gender": "male", "date_of_birth": None, "approx_age_value": 32, "approx_age_unit": "year", "phone": "07701234567", "address": "الموصل"}, self.user)
        imported.source_type = "excel"
        imported.external_id = "XLS-TEST-1"
        imported.save(update_fields=["source_type", "external_id", "updated_at"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("patients:list"), {"phone": "07701234567", "age": "32"})
        self.assertContains(response, imported.internal_code)
        self.assertNotContains(response, manual.internal_code)
        self.assertContains(response, "السجل المستورد")

        response = self.client.get(reverse("patients:match"), {"name": "حيدر سالم", "phone": "07701234567", "age": "32"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["results"]), 1)
        self.assertEqual(response.json()["results"][0]["id"], str(imported.pk))
        self.assertEqual(response.json()["results"][0]["source"], "مستورد")

    def test_registration_allows_two_names_and_optional_contact_address_diagnosis(self):
        form = PatientForm(data={
            "full_name": "علي حسن", "gender": "male",
            "approx_age_value": "30", "approx_age_unit": "year",
            "phone": "", "address": "", "diagnosis_reference": "",
            "department": "", "doctor_reference": "", "organizer_reference": "",
            "visit_date": timezone.localdate().isoformat(),
        }, require_complete=True)
        form.is_valid()
        self.assertNotIn("full_name", form.errors)
        self.assertNotIn("phone", form.errors)
        self.assertNotIn("address", form.errors)
        self.assertNotIn("diagnosis_reference", form.errors)

    def test_doctor_choices_depend_on_department(self):
        women = Department.objects.create(name="نسائية", code="WOMEN")
        children = Department.objects.create(name="اطفال", code="CHILDREN")
        doctor_women = ReferenceValue.objects.create(category="doctor", canonical_name="د. طبيبة نسائية", normalized_name="د طبيبه نسائيه")
        doctor_children = ReferenceValue.objects.create(category="doctor", canonical_name="د. طبيب أطفال", normalized_name="د طبيب اطفال")
        doctor_women.departments.add(women)
        doctor_children.departments.add(children)
        form = PatientForm(department=women)
        self.assertEqual(list(form.fields["doctor_reference"].queryset), [doctor_women])

    def test_patient_drawer_uses_same_screen_data(self):
        patient = create_patient({"full_name": "سارة أحمد محمود علي", "gender": "female", "date_of_birth": None, "approx_age_value": 8, "approx_age_unit": "year", "phone": "07899189225", "address": "الزهور"}, self.user)
        self.client.force_login(self.user)
        response = self.client.get(reverse("patients:drawer", args=[patient.pk]))
        self.assertContains(response, "بطاقة المريض")
        self.assertContains(response, patient.internal_code)

    def test_patient_drawer_saves_without_server_error(self):
        patient = create_patient({"full_name": "سارة أحمد محمود علي", "gender": "female", "date_of_birth": None, "approx_age_value": 8, "approx_age_unit": "year", "phone": "07899189225", "address": "الزهور"}, self.user)
        self.client.force_login(self.user)
        response = self.client.post(reverse("patients:drawer", args=[patient.pk]), {
            "full_name": "سارة أحمد محمود حسن", "gender": "female",
            "approx_age_value": "9", "approx_age_unit": "year",
            "phone": "07899189225", "address": "الزهور",
        }, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        patient.refresh_from_db()
        self.assertEqual(patient.display_name, "سارة أحمد محمود حسن")

    def test_live_search_from_first_character_name_phone_and_id(self):
        patient = create_patient({"full_name": "محمد أحمد علي حسن", "gender": "male", "date_of_birth": None, "approx_age_value": 30, "approx_age_unit": "year", "phone": "07701234567", "address": "الموصل"}, self.user)
        self.client.force_login(self.user)
        for query in ("م", "محمد", "701234", patient.internal_code):
            response = self.client.get(reverse("patients:search"), {"q": query})
            self.assertEqual(response.status_code, 200)
            self.assertIn(str(patient.pk), [item["id"] for item in response.json()["results"]])

    def test_duplicate_candidates_are_safe_and_confirmed_patient_gets_visit(self):
        existing = create_patient({"full_name": "محمد أحمد علي حسن", "gender": "male", "date_of_birth": None, "approx_age_value": 30, "approx_age_unit": "year", "phone": "07701234567", "address": "الموصل"}, self.user)
        self.assertFalse(find_patient_candidates(name="شخص جديد تماماً", phone="07800000000"))
        self.assertEqual(find_patient_candidates(name="محمد أحمد علي حسن", phone="07701234567")[0]["patient"], existing)
        uncertain = find_patient_candidates(name="محمد أحمد علي حسن", phone="07899999999", gender="female", address="البصرة")
        self.assertEqual(uncertain[0]["reasons"], ["الاسم مطابق"])
        self.assertLess(uncertain[0]["score"], 50)  # مرشح بشري فقط، وليس دمجاً تلقائياً

        self.client.force_login(self.user)
        before = Patient.objects.count()
        visit_url = reverse("visits:create_for_patient", args=[existing.pk])
        response = self.client.post(visit_url, {
            "patient_code": existing.internal_code,
            "visit_date": "2026-09-16T10:00", "visit_type": "follow_up", "status": "open",
            "chief_complaint": "متابعة", "diagnosis": "فحص دوري", "notes": "",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Patient.objects.count(), before)
        self.assertTrue(Visit.objects.filter(patient=existing, chief_complaint="متابعة").exists())

    def test_notifications_open_and_mark_all_read_without_deleting(self):
        patient = create_patient({"full_name": "اختبار إشعار مريض علي", "gender": "male", "date_of_birth": None, "approx_age_value": 20, "approx_age_unit": "year", "phone": "07701111111", "address": "الموصل"}, self.user)
        first = Notification.objects.create(user=self.user, title="الأول", target_url=reverse("patients:list") + f"?patient={patient.pk}")
        Notification.objects.create(user=self.user, title="الثاني")
        self.client.force_login(self.user)
        response = self.client.get(reverse("core:notification_open", args=[first.pk]), HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.json()["unread"], 1)
        first.refresh_from_db(); self.assertIsNotNone(first.read_at)
        response = self.client.post(reverse("core:notifications"), HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.json()["unread"], 0)
        self.assertEqual(Notification.objects.count(), 2)

    def test_reference_combo_quick_create_and_reuse(self):
        self.client.force_login(self.user)
        url = reverse("core:reference_quick_create")
        first = self.client.post(url, {"category": "organizer", "name": "يحيى أحمد"}, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(first.status_code, 200)
        self.assertTrue(first.json()["created"])
        second = self.client.post(url, {"category": "organizer", "name": "يحيى أحمد"}, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(second.status_code, 200)
        self.assertFalse(second.json()["created"])
        self.assertEqual(ReferenceValue.objects.filter(category="organizer").count(), 1)

    def test_referral_can_be_edited_with_reference_doctor(self):
        from apps.referrals.models import Referral
        patient = create_patient({"full_name": "مريض إحالة اختبار كامل", "gender": "male", "date_of_birth": None, "approx_age_value": 25, "approx_age_unit": "year", "phone": "07702222222", "address": "الموصل"}, self.user)
        doctor = ReferenceValue.objects.create(category="doctor", canonical_name="طبيب إحالة", normalized_name="طبيب احاله")
        item = Referral.objects.create(patient=patient, destination_name="مستشفى", referral_date=timezone.now())
        self.client.force_login(self.user)
        response = self.client.post(reverse("referrals:edit", args=[item.pk]), {
            "patient_code": patient.internal_code, "referring_doctor_reference": doctor.pk,
            "destination_name": "مستشفى تخصصي", "destination_type": "مستشفى",
            "reason": "استشارة", "referral_date": "2026-09-17T10:00", "status": "pending", "followup_notes": "",
        })
        self.assertEqual(response.status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.referring_doctor_reference, doctor)

    def test_lab_order_uses_shared_reference_doctor(self):
        patient = create_patient({"full_name": "مريض مختبر اختبار كامل", "gender": "male", "date_of_birth": None, "approx_age_value": 25, "approx_age_unit": "year", "phone": "07703333333", "address": "الموصل"}, self.user)
        doctor = ReferenceValue.objects.create(category="doctor", canonical_name="طبيب مختبر حقيقي", normalized_name="طبيب مختبر حقيقي")
        self.client.force_login(self.user)
        response = self.client.post(reverse("laboratory:create_for_patient", args=[patient.pk]), {
            "patient_code": patient.internal_code,
            "order_date": "2026-09-17T10:00",
            "requesting_doctor_reference": doctor.pk,
            "status": "pending",
            "test_name": "CBC",
            "result_value": "",
            "unit": "",
            "notes": "",
        })
        self.assertEqual(response.status_code, 302)
        order = LabOrder.objects.get(patient=patient)
        self.assertEqual(order.requesting_doctor_reference, doctor)
        self.assertEqual(str(order.requesting_doctor_display), "طبيب مختبر حقيقي")

    def test_owner_can_archive_and_restore_reference_without_deleting_history(self):
        owner_role = Role.objects.create(name="مالك القوائم", code=Role.CODE_OWNER)
        owner = User.objects.create_user(username="reference-owner", password="StrongPass123", role=owner_role)
        item = ReferenceValue.objects.create(category="doctor", canonical_name="طبيب قابل للأرشفة", normalized_name="طبيب قابل للارشفة")
        self.client.force_login(owner)
        response = self.client.post(reverse("core:reference_archive", args=[item.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(ReferenceValue.objects.filter(pk=item.pk).exists())
        self.assertTrue(ReferenceValue.all_objects.filter(pk=item.pk, deleted_at__isnull=False).exists())
        response = self.client.post(reverse("core:reference_restore", args=[item.pk]))
        self.assertEqual(response.status_code, 302)
        item.refresh_from_db()
        self.assertTrue(item.is_active)
        self.assertIsNone(item.deleted_at)


class ImportAnalysisTests(TestCase):
    def _workbook(self, path):
        book = Workbook()
        review = book.active; review.title = "مراجعة المرضى"
        review.append(["عنوان"]); review.append([]); review.append(["ت", "الاسم", "الجنس", "العمر", "العنوان", "رقم الهاتف", "عدد التكرار", "القسم", "الحالة", "اسم الطبيب", "اسم المنظم", "التاريخ", "الملاحظات"])
        review.append([1, "علي حسن كامل", "ذكر", 30, "الموصل", "07701234567", 2, "عام", "مراجعة", "د. أ", "منظم", "2026-01-01", ""])
        review.append([2, "علي حسن كامل", "ذكر", 31, "الموصل", "07701234567", 2, "عام", "متابعة", "د. أ", "منظم", "2026-02-01", ""])
        review.append(["=ROW()-3", None, None, None, None, None, "=COUNTIF(B:B,B6)"])
        lab = book.create_sheet("المختبر"); lab.append(["مختبر"]); lab.append(["الاسم", None, "العمر", "نوع الفحص"]); lab.append([1, "نور سامي", 22, "CBC"])
        referrals = book.create_sheet("احالات"); referrals.append(["عنوان"]); referrals.append(["الاسم", None, "الطبيب"]); referrals.append([1, "هدى ياسين", "د. عيون"])
        eye = book.create_sheet("عيادة العيون"); eye.append(["الاسم", None, "العمر", "الجنس", "المنطقة"]); eye.append([1, "سعد وليد", 9, "ذكر", "الرفاعي"])
        book.save(path)

    def test_four_sheet_profiles_skip_formula_counter_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "patients.xlsx"; self._workbook(path)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            batch, created = analyze_workbook(path, digest, path.name, None)
        self.assertTrue(created)
        self.assertEqual(batch.total_rows, 5)
        self.assertEqual(list(batch.sheets.values_list("actual_data_rows", flat=True)), [2, 1, 1, 1])
        self.assertEqual(batch.rows.filter(classification="repeat_visit").count(), 1)
        self.assertTrue(ReferenceValue.objects.filter(category="department", canonical_name="عام").exists())
        self.assertTrue(ReferenceValue.objects.filter(category="lab_test", canonical_name="CBC").exists())
        self.assertTrue(ReferenceValue.objects.filter(category="referral_destination").exists())
        self.assertEqual(batch.rows.filter(raw_data__source__isnull=False).count(), 5)
        imported = import_batch_records(batch, None)
        self.assertGreaterEqual(imported, 1)
        imported_patient = Patient.objects.filter(source_type="excel").first()
        self.assertIsNotNone(imported_patient)
        self.assertEqual(imported_patient.source_file, "patients.xlsx")
        self.assertIn("source_columns", imported_patient.additional_data)
        self.assertEqual(Patient.objects.filter(names__full_name="علي حسن كامل").count(), 1)
        self.assertEqual(Patient.objects.get(names__full_name="علي حسن كامل").visits.count(), 2)
        imported_patient = Patient.objects.get(names__full_name="علي حسن كامل")
        self.assertEqual(imported_patient.imported_visit_count, 2)
        self.assertEqual(imported_patient.total_visit_count, 2)
        self.assertFalse(batch.rows.filter(linked_patient__isnull=True).exists())

    def test_formula_text_never_becomes_visit_count(self):
        raw = {
            "source": {"عدد التكرار": "=COUNTIF($B:$B,B29311)"},
            "evaluated": {"عدد التكرار": 3},
            "canonical": {"repeat_count": "=COUNTIF($B:$B,B29311)"},
        }
        self.assertEqual(source_repeat_count(raw), 3)

    def test_repair_command_corrects_corrupt_count_from_cached_excel_value(self):
        role = Role.objects.create(name="مدقق إصلاح", code=Role.CODE_AUDITOR)
        user = User.objects.create_user(username="repair-auditor", password="StrongPass123", role=role)
        patient = create_patient({"full_name": "مريض إصلاح عداد", "gender": "male", "approx_age_value": 30, "approx_age_unit": "year", "phone": "", "address": "الموصل", "source_type": "excel", "imported_visit_count": 29311}, user)
        batch = ImportBatch.objects.create(file_hash="a" * 64, original_filename="repair.xlsx", import_type="patients", status="completed", imported_by=user)
        sheet = ImportSheet.objects.create(batch=batch, sheet_name="مراجعة المرضى", sheet_index=0)
        SourceRow.objects.create(batch=batch, sheet=sheet, original_row_number=29311, row_hash="b" * 64, linked_patient=patient, status="accepted", raw_data={"source": {"عدد التكرار": "=COUNTIF($B:$B,B29311)"}, "evaluated": {"عدد التكرار": 2}, "canonical": {"name": "مريض إصلاح عداد", "repeat_count": "=COUNTIF($B:$B,B29311)"}})
        call_command("repair_imported_data", verbosity=0)
        patient.refresh_from_db()
        self.assertEqual(patient.imported_visit_count, 2)

    def test_imported_registry_opens_every_imported_patient_drawer(self):
        role = Role.objects.create(name="مدقق", code=Role.CODE_AUDITOR)
        user = User.objects.create_user(username="auditor", password="StrongPass123", role=role)
        patient = create_patient({"full_name": "مريض مستورد كامل علي", "gender": "male", "approx_age_value": 41, "approx_age_unit": "year", "phone": "07701234568", "address": "الموصل"}, user)
        patient.source_type = "excel"
        patient.source_sheet = "مراجعة المرضى"
        patient.source_row = 51
        patient.save(update_fields=["source_type", "source_sheet", "source_row", "updated_at"])
        self.client.force_login(user)
        response = self.client.get(reverse("patients:imported"))
        self.assertContains(response, patient.internal_code)
        response = self.client.get(reverse("patients:drawer", args=[patient.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "بطاقة المريض")

    def test_import_row_detail_with_missing_diagnosis_never_returns_500(self):
        role = Role.objects.create(name="مدقق بطاقات", code=Role.CODE_AUDITOR)
        user = User.objects.create_user(username="row-auditor", password="StrongPass123", role=role)
        batch = ImportBatch.objects.create(file_hash="c" * 64, original_filename="rows.xlsx", import_type="patients", imported_by=user)
        sheet = ImportSheet.objects.create(batch=batch, sheet_name="عيادة العيون", sheet_index=0)
        row = SourceRow.objects.create(
            batch=batch, sheet=sheet, original_row_number=2, row_hash="d" * 64,
            classification="blocking", normalized_name="وسيم انس حسني",
            raw_data={"source": {"الاسم": "وسيم انس حسني"}, "canonical": {"name": "وسيم انس حسني", "age": 9, "gender": "ذكر", "address": "الرفاعي"}},
        )
        self.client.force_login(user)
        response = self.client.get(reverse("importer:row_detail", args=[row.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "وسيم انس حسني")
        self.assertContains(response, "إضافة إلى السجل المستورد والمرضى")

    def test_blocking_duplicate_and_repeat_rows_are_all_linked(self):
        role = Role.objects.create(name="مدقق كل التصنيفات", code=Role.CODE_AUDITOR)
        user = User.objects.create_user(username="all-rows-auditor", password="StrongPass123", role=role)
        batch = ImportBatch.objects.create(file_hash="e" * 64, original_filename="all.xlsx", import_type="patients", imported_by=user)
        sheet = ImportSheet.objects.create(batch=batch, sheet_name="مراجعة المرضى", sheet_index=0)
        for index, classification in enumerate(("ready", "review", "blocking", "duplicate", "repeat_visit"), 2):
            SourceRow.objects.create(
                batch=batch, sheet=sheet, original_row_number=index,
                row_hash=(str(index) * 64)[:64], classification=classification,
                normalized_name="مريض واحد كامل علي",
                raw_data={"source": {"الاسم": "مريض واحد كامل علي"}, "canonical": {"name": "مريض واحد كامل علي", "gender": "ذكر"}},
            )
        imported = import_batch_records(batch, user)
        self.assertEqual(imported, 5)
        self.assertFalse(batch.rows.filter(linked_patient__isnull=True).exists())
        self.assertEqual(Patient.objects.filter(names__normalized_name="مريض واحد كامل علي").distinct().count(), 1)

    def test_every_authenticated_page_has_global_back_button(self):
        role = Role.objects.create(name="مدقق زر الرجوع", code=Role.CODE_AUDITOR)
        user = User.objects.create_user(username="back-auditor", password="StrongPass123", role=role)
        self.client.force_login(user)
        response = self.client.get(reverse("importer:list"))
        self.assertContains(response, "data-global-back")
        self.assertContains(response, "رجوع")

    def test_gender_variants_are_normalized_without_blocking(self):
        from apps.importer.services import GENDERS, normalize_arabic

        for raw, expected in (("رجل", "male"), ("Male", "male"), ("امرأة", "female"), ("F", "female")):
            self.assertEqual(GENDERS[normalize_arabic(raw)], expected)


class BackupExportTests(TestCase):
    def setUp(self):
        role = Role.objects.create(name="المالك", code=Role.CODE_OWNER)
        self.user = User.objects.create_user(username="owner-export", password="StrongPass123", role=role)
        self.patient = create_patient({"full_name": "مريض تصدير كامل علي", "gender": "male", "approx_age_value": 35, "approx_age_unit": "year", "phone": "07705555555", "address": "الموصل"}, self.user)
        self.client.force_login(self.user)

    def _download(self, route):
        response = self.client.get(reverse(route))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        return b"".join(response.streaming_content)

    def test_excel_export_uses_requested_names_columns_and_filterable_tables(self):
        with tempfile.TemporaryDirectory() as directory, override_settings(DATA_PATH=Path(directory)):
            excel = self._download("backup:excel_export")
            workbook = load_workbook(io.BytesIO(excel))
            patients = workbook["المرضى"]
            self.assertEqual([cell.value for cell in patients[1]], [
                "ت", "الرقم التعريفي الخاص بالمريض", "الاسم", "الجنس", "العمر", "العنوان",
                "رقم الهاتف", "القسم", "الحالة", "اسم الطبيب", "اسم المنظم", "التاريخ", "الملاحظات",
            ])
            self.assertEqual(patients["B2"].value, self.patient.internal_code)
            self.assertEqual(patients["C2"].value, self.patient.display_name)
            self.assertTrue(patients.tables)
            self.assertTrue(workbook["المختبر"].tables or workbook["المختبر"].max_row == 1)
            workbook.close()

    def test_excel_csv_and_json_are_real_downloads(self):
        with tempfile.TemporaryDirectory() as directory, override_settings(DATA_PATH=Path(directory)):
            excel = self._download("backup:excel_export")
            workbook = load_workbook(io.BytesIO(excel), read_only=True)
            self.assertIn("المرضى", workbook.sheetnames)
            self.assertGreaterEqual(sum(1 for _ in workbook["المرضى"].iter_rows(values_only=True)), 2)
            workbook.close()

            csv_archive = self._download("backup:csv_export")
            with zipfile.ZipFile(io.BytesIO(csv_archive)) as archive:
                self.assertIn("المرضى.csv", archive.namelist())

            json_data = json.loads(self._download("backup:json_export").decode("utf-8"))
            self.assertEqual(json_data["المرضى"][0]["الرقم التعريفي الخاص بالمريض"], self.patient.internal_code)


class ReleaseReadinessTests(TestCase):
    def setUp(self):
        self.role = Role.objects.create(name="مدقق جاهزية", code=Role.CODE_AUDITOR)
        self.user = User.objects.create_user(username="readiness-auditor", password="StrongPass123", role=self.role)
        self.client.force_login(self.user)

    def test_language_switch_persists_arabic_and_english(self):
        response = self.client.get(reverse("core:set_language", args=["ar"]))
        self.assertEqual(response.cookies["django_language"].value, "ar")
        self.assertEqual(self.client.session["django_language"], "ar")
        response = self.client.get(reverse("core:set_language", args=["en"]))
        self.assertEqual(response.cookies["django_language"].value, "en")

    def test_patient_trash_restores_without_permanent_deletion(self):
        patient = create_patient({"full_name": "مريض سلة قابل للاستعادة", "gender": "male", "approx_age_value": 22, "approx_age_unit": "year", "phone": "", "address": ""}, self.user)
        patient.soft_delete()
        response = self.client.get(reverse("patients:trash"))
        self.assertContains(response, patient.internal_code)
        response = self.client.post(reverse("patients:restore", args=[patient.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Patient.objects.filter(pk=patient.pk).exists())

    def test_exports_are_grouped_in_named_data_folders(self):
        with tempfile.TemporaryDirectory() as directory, override_settings(DATA_PATH=Path(directory)):
            from apps.backup.services import create_excel_export, create_csv_export, create_json_export
            self.assertEqual(create_excel_export().parent.name, "Excel")
            self.assertEqual(create_csv_export().parent.name, "ZIP")
            self.assertEqual(create_json_export().parent.name, "JSON")
